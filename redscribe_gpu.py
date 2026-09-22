import tkinter as tk
from tkinter import filedialog, messagebox, scrolledtext, ttk
from faster_whisper import WhisperModel
import threading
import time
import os
import re
import shutil
import subprocess
import tempfile
from collections import Counter

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList


# =========================
# CONFIG
# =========================
MODEL_SIZE = "large-v3"

# Whisper priority:
# cuba GPU dulu, kalau fail baru fallback CPU
WHISPER_DEVICE_PRIORITY = ["cuda", "cpu"]

# compute type ikut device
WHISPER_COMPUTE_TYPE = {
    "cuda": "float16",   # laju untuk GPU NVIDIA
    "cpu": "int8"        # ringan untuk CPU
}

# Diarization model.
# Nota:
# 1. Install: pip install pyannote.audio
# 2. Login Hugging Face atau letak token dalam GUI / environment variable HF_TOKEN.
# 3. Accept model terms dekat Hugging Face kalau diminta.
DIARIZATION_MODEL = "pyannote/speaker-diarization-community-1"

APP_VERSION = "2.0.0"
APP_TITLE = f"RedScribe v{APP_VERSION} – Research Speech Transcription System"


# =========================
# UTILS
# =========================
def sec_to_hms(sec: float) -> str:
    sec = max(0, int(sec))
    h = sec // 3600
    m = (sec % 3600) // 60
    s = sec % 60
    return f"{h:02d}:{m:02d}:{s:02d}"


def sec_to_mmss(sec: float) -> str:
    sec = max(0, int(sec))
    m = sec // 60
    s = sec % 60
    return f"{m:02d}:{s:02d}"


def split_sentences_universal(text: str) -> list[str]:
    text = (text or "").strip()
    if not text:
        return []
    parts = re.split(r'(?<=[\.\!\?\u061F\u3002\uFF01\uFF1F])\s+', text)
    parts = [p.strip() for p in parts if p and p.strip()]
    return parts if parts else [text]


def distribute_times(start: float, end: float, sentences: list[str]) -> list[tuple[float, float, str]]:
    dur = max(0.0, end - start)
    if not sentences:
        return []
    if dur == 0:
        return [(start, end, s) for s in sentences]

    weights = [max(1, len(s)) for s in sentences]
    total = float(sum(weights))
    out = []
    t = float(start)

    for i, s in enumerate(sentences):
        if i == len(sentences) - 1:
            t2 = float(end)
        else:
            t2 = t + dur * (weights[i] / total)
        out.append((t, t2, s))
        t = t2
    return out


def basic_tokenize(text: str) -> list[str]:
    text = (text or "").lower()
    tokens = re.findall(r"[a-z0-9]+", text)
    return [t for t in tokens if len(t) >= 3]


def compute_timeline_density(utterance_rows: list[dict], audio_duration_s: float, interval_sec: int = 300):
    buckets = Counter()
    for r in utterance_rows:
        b = int(r["start_s"] // interval_sec)
        buckets[b] += 1

    max_bucket = int(audio_duration_s // interval_sec) if audio_duration_s > 0 else 0
    rows = []
    for b in range(0, max_bucket + 1):
        start_min = b * (interval_sec // 60)
        end_min = start_min + (interval_sec // 60)
        rows.append((f"{start_min}-{end_min}", buckets.get(b, 0)))
    return rows


def compute_speaker_summary(utterance_rows: list[dict]) -> list[tuple[str, int]]:
    speaker_counts = Counter()
    for row in utterance_rows:
        speaker = (row.get("speaker") or "UNKNOWN").strip()
        speaker_counts[speaker] += 1
    return speaker_counts.most_common()


def sanitize_part(s: str) -> str:
    s = (s or "").strip()
    if not s:
        return ""
    s = re.sub(r"\s+", " ", s)
    s = re.sub(r'[<>:"/\\|?*\n\r\t]', "_", s)
    s = s.replace(" ", "_")
    s = re.sub(r"_+", "_", s)
    return s.strip("_")


def build_output_name(project: str, period: str, label: str, researcher: str, audio_base: str) -> str:
    parts = [
        sanitize_part(project),
        sanitize_part(period),
        sanitize_part(label),
        sanitize_part(researcher),
    ]
    parts = [p for p in parts if p]
    if not parts:
        return sanitize_part(audio_base) or "transcription"
    return "_".join(parts)


def ensure_dir(path: str):
    os.makedirs(path, exist_ok=True)


def open_folder(path: str):
    try:
        os.startfile(path)
    except Exception:
        pass


def create_whisper_model():
    """
    Cuba load whisper guna GPU dulu.
    Kalau gagal, fallback ke CPU.
    Return:
        model, actual_device, actual_compute_type
    """
    last_error = None

    for device in WHISPER_DEVICE_PRIORITY:
        compute_type = WHISPER_COMPUTE_TYPE.get(device, "int8")
        try:
            model = WhisperModel(MODEL_SIZE, device=device, compute_type=compute_type)
            return model, device, compute_type
        except Exception as e:
            last_error = e

    raise RuntimeError(f"Gagal load Whisper model pada semua device. Error terakhir: {last_error}")



def load_audio_for_pyannote(audio_path: str):
    """
    Load audio as an in-memory waveform for pyannote.

    Why this exists:
    Some Windows setups fail inside pyannote/torchcodec with:
    "name 'AudioDecoder' is not defined".
    Passing an already-loaded waveform avoids pyannote's built-in AudioDecoder path.

    Requires:
        pip install soundfile
        ffmpeg available in PATH
    """
    tmp_dir = tempfile.mkdtemp(prefix="redscribe_diarization_")
    wav_path = os.path.join(tmp_dir, "audio_16k_mono.wav")

    try:
        cmd = [
            "ffmpeg",
            "-y",
            "-hide_banner",
            "-loglevel", "error",
            "-i", audio_path,
            "-vn",
            "-ac", "1",
            "-ar", "16000",
            wav_path,
        ]
        try:
            subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, check=True, text=True)
        except FileNotFoundError:
            raise RuntimeError(
                "FFmpeg tidak dijumpai dalam PATH. Install FFmpeg dulu, kemudian buka semula PowerShell. "
                "Contoh: winget install Gyan.FFmpeg"
            )
        except subprocess.CalledProcessError as e:
            err = (e.stderr or e.stdout or "").strip()
            raise RuntimeError(f"FFmpeg gagal convert audio untuk diarization. Error: {err}")

        try:
            import soundfile as sf
            import torch
        except Exception as e:
            raise RuntimeError(
                "Package audio loader belum lengkap. Install dulu dengan: pip install soundfile\n\n"
                f"Error asal: {e}"
            )

        data, sample_rate = sf.read(wav_path, dtype="float32", always_2d=False)
        if getattr(data, "ndim", 1) > 1:
            data = data.mean(axis=1)

        waveform = torch.from_numpy(data).float().unsqueeze(0)
        return {"waveform": waveform, "sample_rate": int(sample_rate)}

    finally:
        try:
            shutil.rmtree(tmp_dir, ignore_errors=True)
        except Exception:
            pass


def create_diarization_pipeline(hf_token: str | None, preferred_device: str = "cpu"):
    """
    Load pyannote diarization pipeline.
    Function ni lazy import supaya app masih boleh buka walaupun pyannote.audio belum install.
    """
    try:
        from pyannote.audio import Pipeline
    except Exception as e:
        raise RuntimeError(
            "pyannote.audio belum dipasang. Install dulu dengan:\n"
            "pip install pyannote.audio\n\n"
            f"Error asal: {e}"
        )

    token = (
        (hf_token or "").strip()
        or os.environ.get("HF_TOKEN")
        or os.environ.get("HUGGINGFACE_HUB_TOKEN")
        or None
    )

    try:
        if token:
            pipeline = Pipeline.from_pretrained(DIARIZATION_MODEL, token=token)
        else:
            # token=True cuba guna cached login daripada huggingface-cli login
            pipeline = Pipeline.from_pretrained(DIARIZATION_MODEL, token=True)
    except TypeError:
        # fallback untuk versi lama pyannote / huggingface_hub
        if token:
            pipeline = Pipeline.from_pretrained(DIARIZATION_MODEL, use_auth_token=token)
        else:
            pipeline = Pipeline.from_pretrained(DIARIZATION_MODEL, use_auth_token=True)

    if pipeline is None:
        raise RuntimeError(
            "Diarization model gagal dimuatkan. Pastikan Hugging Face token betul "
            "dan model terms sudah diterima di Hugging Face."
        )

    # Hantar pyannote ke GPU jika ada dan Whisper berjaya guna CUDA.
    if preferred_device == "cuda":
        try:
            import torch
            if torch.cuda.is_available():
                pipeline.to(torch.device("cuda"))
        except Exception:
            pass

    return pipeline


def run_diarization(audio_path: str, hf_token: str | None, preferred_device: str, speaker_count_text: str = "") -> list[dict]:
    """
    Return list:
    [
        {"start": 0.0, "end": 3.2, "speaker": "SPEAKER_00"},
        ...
    ]
    """
    pipeline = create_diarization_pipeline(hf_token, preferred_device)

    kwargs = {}
    speaker_count_text = (speaker_count_text or "").strip()
    if speaker_count_text:
        try:
            n_speakers = int(speaker_count_text)
            if n_speakers > 0:
                kwargs["num_speakers"] = n_speakers
        except ValueError:
            pass

    diarization_input = load_audio_for_pyannote(audio_path)
    diarization_output = pipeline(diarization_input, **kwargs)

    # pyannote/speaker-diarization-community-1 returns a DiarizeOutput object.
    # The actual Annotation is inside .exclusive_speaker_diarization or .speaker_diarization.
    # Older pyannote pipelines may return the Annotation directly.
    diarization_annotation = None
    for attr_name in ("exclusive_speaker_diarization", "speaker_diarization"):
        if hasattr(diarization_output, attr_name):
            diarization_annotation = getattr(diarization_output, attr_name)
            if diarization_annotation is not None:
                break

    if diarization_annotation is None:
        diarization_annotation = diarization_output

    if not hasattr(diarization_annotation, "itertracks"):
        raise RuntimeError(
            "Pyannote output tidak mempunyai itertracks(). "
            f"Output type: {type(diarization_output).__name__}. "
            "Sila pastikan pyannote.audio versi terbaru dan guna model community-1."
        )

    turns = []
    for turn, _, speaker in diarization_annotation.itertracks(yield_label=True):
        turns.append({
            "start": float(turn.start),
            "end": float(turn.end),
            "speaker": str(speaker)
        })

    turns.sort(key=lambda x: (x["start"], x["end"]))
    return turns


def overlap_seconds(a_start: float, a_end: float, b_start: float, b_end: float) -> float:
    return max(0.0, min(a_end, b_end) - max(a_start, b_start))


def assign_speaker(start: float, end: float, diarization_turns: list[dict]) -> str:
    """
    Assign speaker based on maximum overlap between Whisper timestamp and pyannote diarization turns.
    Kalau tiada overlap, guna midpoint fallback.
    """
    if not diarization_turns:
        return "UNKNOWN"

    best_speaker = "UNKNOWN"
    best_overlap = 0.0

    for turn in diarization_turns:
        ov = overlap_seconds(start, end, turn["start"], turn["end"])
        if ov > best_overlap:
            best_overlap = ov
            best_speaker = turn["speaker"]

    if best_overlap > 0:
        return best_speaker

    midpoint = (start + end) / 2.0
    for turn in diarization_turns:
        if turn["start"] <= midpoint <= turn["end"]:
            return turn["speaker"]

    return "UNKNOWN"


# =========================
# EXCEL WRITER
# =========================
def save_xlsx_with_early_report(
    xlsx_path: str,
    utterance_rows: list[dict],
    segment_ranges: list[tuple[float, float]],
    audio_path: str,
    audio_duration_s: float,
    processing_minutes: float,
    detected_language: str | None,
    meta_from_gui: dict,
    actual_device: str,
    actual_compute_type: str,
    diarization_enabled: bool,
    diarization_status: str,
    diarization_model: str
):
    wb = Workbook()

    ws1 = wb.active
    ws1.title = "Early_Report"

    ws1["A1"] = "Study & Processing Metadata"
    ws1["A1"].font = ws1["A1"].font.copy(bold=True)

    speaker_summary = compute_speaker_summary(utterance_rows)
    speakers_detected = ", ".join([speaker for speaker, _ in speaker_summary]) if speaker_summary else "None"

    meta = [
        ("Application", f"RedScribe v{APP_VERSION} – Research Speech Transcription System"),
        ("Project / Study", meta_from_gui.get("project", "")),
        ("Researcher", meta_from_gui.get("researcher", "")),
        ("Interview / Audio Label", meta_from_gui.get("label", "")),
        ("Processing Period", meta_from_gui.get("period", "")),
        ("Audio File", os.path.basename(audio_path)),
        ("Audio Duration", sec_to_hms(audio_duration_s)),
        ("Language (Detected/Selected)", detected_language or "Auto detect"),
        ("Model & Mode", f"{MODEL_SIZE} ({actual_device.upper()} | {actual_compute_type})"),
        ("Diarization", "Enabled" if diarization_enabled else "Disabled"),
        ("Diarization Status", diarization_status),
        ("Diarization Model", diarization_model if diarization_enabled else "Not used"),
        ("Speakers Detected", speakers_detected),
        ("Processing Time", f"{processing_minutes} minutes"),
        ("Generated On", time.strftime("%Y-%m-%d %H:%M:%S")),
        ("Copyright", "© Mad Sapri Tumiran"),
        ("Support", "saprimad@moh.gov.my"),
    ]

    r = 3
    for k, v in meta:
        ws1[f"A{r}"] = k
        ws1[f"B{r}"] = v
        r += 1

    ws1[f"A{r+1}"] = "Transcription Summary"
    ws1[f"A{r+1}"].font = ws1[f"A{r+1}"].font.copy(bold=True)

    total_segments = len(segment_ranges)
    total_utterances = len(utterance_rows)

    all_text = " ".join([x["text"] for x in utterance_rows])
    tokens = basic_tokenize(all_text)
    token_count = len(tokens)
    avg_tokens = round(token_count / total_utterances, 2) if total_utterances else 0.0

    speech_s = 0.0
    for st, en in segment_ranges:
        if en > st:
            speech_s += (en - st)
    speech_s = max(0.0, min(speech_s, audio_duration_s))
    silence_s = max(0.0, audio_duration_s - speech_s)

    summary = [
        ("Total Segments Detected", total_segments),
        ("Total Utterances", total_utterances),
        ("Total Words/Tokens", token_count),
        ("Average Tokens per Utterance", avg_tokens),
        ("Speech Duration", sec_to_hms(speech_s)),
        ("Silence Duration", sec_to_hms(silence_s)),
        ("Total Speakers Detected", len([s for s, _ in speaker_summary if s != "UNKNOWN"])),
    ]

    r2 = r + 3
    for k, v in summary:
        ws1[f"A{r2}"] = k
        ws1[f"B{r2}"] = v
        r2 += 1

    r_spk = r2 + 2
    ws1[f"A{r_spk}"] = "Speaker Utterance Count"
    ws1[f"A{r_spk}"].font = ws1[f"A{r_spk}"].font.copy(bold=True)
    ws1[f"A{r_spk+2}"] = "Speaker"
    ws1[f"B{r_spk+2}"] = "Utterances"

    rr_spk = r_spk + 3
    for speaker, count in speaker_summary:
        ws1[f"A{rr_spk}"] = speaker
        ws1[f"B{rr_spk}"] = count
        rr_spk += 1

    r3 = max(rr_spk + 2, r2 + 5)
    ws1[f"A{r3}"] = "Speech vs Silence Distribution"
    ws1[f"A{r3}"].font = ws1[f"A{r3}"].font.copy(bold=True)

    ws1[f"A{r3+2}"] = "Category"
    ws1[f"B{r3+2}"] = "Duration (minutes)"

    speech_min = round(speech_s / 60.0, 3)
    silence_min = round(silence_s / 60.0, 3)

    ws1[f"A{r3+3}"] = "Speech"
    ws1[f"B{r3+3}"] = speech_min
    ws1[f"A{r3+4}"] = "Silence"
    ws1[f"B{r3+4}"] = silence_min

    chart1 = BarChart()
    chart1.type = "col"
    chart1.title = "Speech vs Silence Distribution"
    chart1.y_axis.title = "Duration (minutes)"
    chart1.x_axis.title = "Category"
    chart1.dataLabels = DataLabelList()
    chart1.dataLabels.showVal = True

    data = Reference(ws1, min_col=2, min_row=r3+2, max_row=r3+4)
    cats = Reference(ws1, min_col=1, min_row=r3+3, max_row=r3+4)
    chart1.add_data(data, titles_from_data=True)
    chart1.set_categories(cats)
    ws1.add_chart(chart1, f"D{r3+2}")

    r4 = r3 + 7
    ws1[f"A{r4}"] = "Timeline Density (5-minute intervals)"
    ws1[f"A{r4}"].font = ws1[f"A{r4}"].font.copy(bold=True)

    ws1[f"A{r4+2}"] = "Interval (min)"
    ws1[f"B{r4+2}"] = "Utterances"

    density_rows = compute_timeline_density(utterance_rows, audio_duration_s, interval_sec=300)
    start_table_row = r4 + 3
    rr = start_table_row
    for interval_label, count in density_rows:
        ws1[f"A{rr}"] = interval_label
        ws1[f"B{rr}"] = count
        rr += 1
    end_table_row = rr - 1

    chart2 = BarChart()
    chart2.type = "col"
    chart2.title = "Utterance Density Over Time"
    chart2.y_axis.title = "Utterance count"
    chart2.x_axis.title = "Time interval (min)"
    chart2.dataLabels = DataLabelList()
    chart2.dataLabels.showVal = False

    data2 = Reference(ws1, min_col=2, min_row=r4+2, max_row=end_table_row)
    cats2 = Reference(ws1, min_col=1, min_row=r4+3, max_row=end_table_row)
    chart2.add_data(data2, titles_from_data=True)
    chart2.set_categories(cats2)
    ws1.add_chart(chart2, f"D{r4+2}")

    r5 = end_table_row + 3
    ws1[f"A{r5}"] = "Top 20 Keywords (Frequency)"
    ws1[f"A{r5}"].font = ws1[f"A{r5}"].font.copy(bold=True)

    ws1[f"A{r5+2}"] = "Rank"
    ws1[f"B{r5+2}"] = "Keyword"
    ws1[f"C{r5+2}"] = "Frequency"

    freq = Counter(tokens).most_common(20)
    rrk = r5 + 3
    for i, (kw, n) in enumerate(freq, start=1):
        ws1[f"A{rrk}"] = i
        ws1[f"B{rrk}"] = kw
        ws1[f"C{rrk}"] = n
        rrk += 1

    ws1.column_dimensions["A"].width = 32
    ws1.column_dimensions["B"].width = 52
    ws1.column_dimensions["C"].width = 16
    ws1.column_dimensions["D"].width = 16
    ws1.column_dimensions["E"].width = 16
    ws1.column_dimensions["F"].width = 16

    ws2 = wb.create_sheet("Transcription")
    headers = [
        "Utterance No.",
        "Speaker",
        "Start Time (hh:mm:ss)",
        "End Time (hh:mm:ss)",
        "Auto Transcript (RedScribe)",
        "Researcher Verification / Correction",
        "Notes"
    ]
    ws2.append(headers)

    for i, rrow in enumerate(utterance_rows, start=1):
        ws2.append([
            i,
            rrow.get("speaker", "UNKNOWN"),
            sec_to_hms(rrow["start_s"]),
            sec_to_hms(rrow["end_s"]),
            rrow["text"],
            "",
            ""
        ])

    widths = [14, 18, 20, 18, 70, 40, 25]
    for col_idx, w in enumerate(widths, start=1):
        ws2.column_dimensions[get_column_letter(col_idx)].width = w

    wb.save(xlsx_path)


# =========================
# MAIN APP
# =========================
class RedScribeApp:
    def __init__(self, root):
        self.root = root
        self.root.title(APP_TITLE)
        self.root.geometry("1280x760")

        self.project = tk.StringVar()
        self.researcher = tk.StringVar()
        self.label = tk.StringVar()
        self.period = tk.StringVar(value=time.strftime("%Y-%m"))
        self.audio_path = tk.StringVar()

        self.diarization_enabled = tk.BooleanVar(value=True)
        self.hf_token = tk.StringVar()
        self.speaker_count = tk.StringVar()

        self.base_output_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), "output")
        self.last_output_folder = None
        self.progress_var = tk.DoubleVar(value=0.0)

        self.actual_whisper_device = "unknown"
        self.actual_whisper_compute_type = "unknown"

        self.build_ui()

    def build_ui(self):
        top = tk.Frame(self.root)
        top.pack(fill="x", padx=10, pady=6)

        tk.Label(top, text="Project & File Information", font=("Segoe UI", 10, "bold")).pack(anchor="w")

        def row_entry(parent, label, var, show=None):
            fr = tk.Frame(parent)
            fr.pack(fill="x", pady=2)
            tk.Label(fr, text=label, width=28, anchor="w").pack(side="left")
            tk.Entry(fr, textvariable=var, show=show).pack(side="left", fill="x", expand=True)
            return fr

        row_entry(top, "Project / Study Name", self.project)
        row_entry(top, "Researcher Name / Initials", self.researcher)
        row_entry(top, "Interview / Audio Label", self.label)

        period_fr = tk.Frame(top)
        period_fr.pack(fill="x", pady=2)
        tk.Label(period_fr, text="Month / Period (YYYY-MM)", width=28, anchor="w").pack(side="left")
        tk.Entry(period_fr, textvariable=self.period, width=12).pack(side="left")
        tk.Button(period_fr, text="Pick…", command=self.pick_period).pack(side="left", padx=6)

        file_fr = tk.Frame(top)
        file_fr.pack(fill="x", pady=2)
        tk.Label(file_fr, text="Audio File", width=28, anchor="w").pack(side="left")
        tk.Entry(file_fr, textvariable=self.audio_path).pack(side="left", padx=4, fill="x", expand=True)
        tk.Button(file_fr, text="Browse…", command=self.browse_file).pack(side="left")

        tk.Label(top, text="Supported formats: mp3, wav, m4a, mp4", fg="gray").pack(anchor="w", pady=(2, 6))

        diar_fr = tk.LabelFrame(top, text="Speaker Diarization")
        diar_fr.pack(fill="x", pady=(4, 6))

        tk.Checkbutton(
            diar_fr,
            text="Enable speaker diarization",
            variable=self.diarization_enabled
        ).pack(anchor="w", padx=6, pady=(4, 2))

        token_fr = tk.Frame(diar_fr)
        token_fr.pack(fill="x", padx=6, pady=2)
        tk.Label(token_fr, text="Hugging Face Token", width=28, anchor="w").pack(side="left")
        tk.Entry(token_fr, textvariable=self.hf_token, show="*").pack(side="left", fill="x", expand=True)

        speaker_fr = tk.Frame(diar_fr)
        speaker_fr.pack(fill="x", padx=6, pady=(2, 6))
        tk.Label(speaker_fr, text="Expected Speakers (optional)", width=28, anchor="w").pack(side="left")
        tk.Entry(speaker_fr, textvariable=self.speaker_count, width=8).pack(side="left")
        tk.Label(
            speaker_fr,
            text="Contoh: 2 untuk interview dua orang. Kosongkan untuk auto detect.",
            fg="gray"
        ).pack(side="left", padx=8)

        main_panel = tk.Frame(self.root)
        main_panel.pack(fill="both", expand=True, padx=10, pady=5)

        tk.Label(main_panel, text="Live Transcript", font=("Segoe UI", 10, "bold")).pack(anchor="w")
        self.live_box = scrolledtext.ScrolledText(main_panel, height=24, state="disabled", wrap="word")
        self.live_box.pack(fill="both", expand=True)

        self.live_box.tag_configure("ts", foreground="blue")
        self.live_box.tag_configure("speaker", foreground="purple")
        self.live_box.tag_configure("txt", foreground="black")
        self.live_box.tag_configure("warn", foreground="red")

        status_row = tk.Frame(self.root)
        status_row.pack(fill="x", padx=10, pady=(0, 6))

        self.status = tk.StringVar(value="Idle")
        tk.Label(status_row, textvariable=self.status).pack(side="left")

        self.progress = ttk.Progressbar(
            status_row,
            orient="horizontal",
            mode="determinate",
            maximum=100,
            variable=self.progress_var,
            length=320
        )
        self.progress.pack(side="right")

        ctrl = tk.Frame(self.root)
        ctrl.pack(fill="x", padx=10, pady=6)

        self.start_btn = tk.Button(
            ctrl,
            text="Start Transcription",
            command=self.start,
            bg="#BFE8BF",
            activebackground="#A7DCA7",
            relief="raised",
            padx=10,
            pady=4
        )
        self.start_btn.pack(side="right", padx=4)

        self.open_btn = tk.Button(
            ctrl,
            text="Open Output Folder",
            command=self.open_output_folder,
            bg="#FFE59A",
            activebackground="#FFD56A",
            relief="raised",
            padx=10,
            pady=4
        )
        self.open_btn.pack(side="right", padx=4)

        footer = tk.Label(self.root, text="© Mad Sapri Tumiran | Support", fg="gray", cursor="hand2")
        footer.pack(side="bottom", pady=3)
        footer.bind("<Button-1>", self.show_support)

    def show_support(self, _event=None):
        messagebox.showinfo(
            "About RedScribe",
            f"RedScribe v{APP_VERSION}\n"
            "Research Speech Transcription System\n\n"
            "© Mad Sapri Tumiran\n\n"
            "Support:\n"
            "saprimad@moh.gov.my"
        )

    def pick_period(self):
        win = tk.Toplevel(self.root)
        win.title("Select Month / Period")
        win.resizable(False, False)
        win.grab_set()

        cur = (self.period.get() or time.strftime("%Y-%m")).strip()
        try:
            y0, m0 = cur.split("-")
            y0 = int(y0)
            m0 = int(m0)
        except Exception:
            y0 = int(time.strftime("%Y"))
            m0 = int(time.strftime("%m"))

        year_var = tk.IntVar(value=y0)
        month_var = tk.IntVar(value=m0)

        tk.Label(win, text="Year").grid(row=0, column=0, padx=10, pady=8, sticky="w")
        tk.Spinbox(win, from_=2000, to=2100, textvariable=year_var, width=8).grid(row=0, column=1, padx=10, pady=8)

        tk.Label(win, text="Month").grid(row=1, column=0, padx=10, pady=8, sticky="w")
        tk.Spinbox(win, from_=1, to=12, textvariable=month_var, width=8).grid(row=1, column=1, padx=10, pady=8)

        def apply():
            y = year_var.get()
            m = month_var.get()
            self.period.set(f"{y:04d}-{m:02d}")
            win.destroy()

        btns = tk.Frame(win)
        btns.grid(row=2, column=0, columnspan=2, pady=10)
        tk.Button(btns, text="Cancel", command=win.destroy).pack(side="right", padx=6)
        tk.Button(btns, text="Apply", command=apply).pack(side="right", padx=6)

    def browse_file(self):
        path = filedialog.askopenfilename(
            filetypes=[("Audio/video files", "*.mp3 *.wav *.m4a *.mp4")]
        )
        if path:
            self.audio_path.set(path)

    def open_output_folder(self):
        if self.last_output_folder and os.path.isdir(self.last_output_folder):
            open_folder(self.last_output_folder)
        else:
            ensure_dir(self.base_output_dir)
            open_folder(self.base_output_dir)

    def log_live_segment(self, ts: str, speaker: str, text: str):
        self.live_box.configure(state="normal")
        self.live_box.insert("end", f"[{ts}] ", ("ts",))
        self.live_box.insert("end", f"{speaker}: ", ("speaker",))
        self.live_box.insert("end", f"{text}\n", ("txt",))
        self.live_box.see("end")
        self.live_box.configure(state="disabled")

    def log_warning(self, text: str):
        self.live_box.configure(state="normal")
        self.live_box.insert("end", f"{text}\n", ("warn",))
        self.live_box.see("end")
        self.live_box.configure(state="disabled")

    def start(self):
        inputs = {
            "audio_path": self.audio_path.get().strip(),
            "project": self.project.get().strip(),
            "researcher": self.researcher.get().strip(),
            "label": self.label.get().strip(),
            "period": self.period.get().strip(),
            "diarization_enabled": bool(self.diarization_enabled.get()),
            "hf_token": self.hf_token.get().strip(),
            "speaker_count": self.speaker_count.get().strip(),
        }

        if not inputs["audio_path"]:
            messagebox.showwarning("Missing file", "Please select an audio file.")
            return

        if not os.path.isfile(inputs["audio_path"]):
            messagebox.showwarning("File not found", "The selected audio file no longer exists.")
            return

        self.live_box.configure(state="normal")
        self.live_box.delete("1.0", "end")
        self.live_box.configure(state="disabled")

        self.progress_var.set(0.0)
        self.status.set("Loading model...")
        self.start_btn.config(state="disabled")

        threading.Thread(target=self.run_whisper, args=(inputs,), daemon=True).start()

    def run_whisper(self, inputs: dict):
        start_clock = time.time()

        audio_path = inputs["audio_path"]
        audio_base = os.path.splitext(os.path.basename(audio_path))[0]

        period = inputs["period"] or time.strftime("%Y-%m")
        if not re.match(r"^\d{4}-\d{2}$", period):
            period = time.strftime("%Y-%m")

        out_folder = os.path.join(self.base_output_dir, period)
        ensure_dir(out_folder)
        self.last_output_folder = out_folder

        filename_base = build_output_name(
            project=inputs["project"],
            period=period,
            label=inputs["label"],
            researcher=inputs["researcher"],
            audio_base=audio_base
        )
        xlsx_path = os.path.join(out_folder, f"{filename_base}.xlsx")

        try:
            model, actual_device, actual_compute_type = create_whisper_model()
            self.actual_whisper_device = actual_device
            self.actual_whisper_compute_type = actual_compute_type
        except Exception as e:
            self.root.after(
                0,
                lambda: messagebox.showerror("Whisper Error", f"Gagal load model Whisper.\n\n{e}")
            )
            self.root.after(0, self.status.set, "Failed to load model")
            self.root.after(0, lambda: self.start_btn.config(state="normal"))
            return

        diarization_enabled = inputs["diarization_enabled"]
        diarization_turns = []
        diarization_status = "Disabled"

        if diarization_enabled:
            self.root.after(0, self.status.set, "Running speaker diarization...")
            try:
                diarization_turns = run_diarization(
                    audio_path=audio_path,
                    hf_token=inputs["hf_token"],
                    preferred_device=actual_device,
                    speaker_count_text=inputs["speaker_count"]
                )
                speakers = sorted({t["speaker"] for t in diarization_turns})
                diarization_status = f"Completed ({len(speakers)} speaker label(s))"
            except Exception as e:
                diarization_status = f"Skipped / failed: {e}"
                self.root.after(
                    0,
                    self.log_warning,
                    "Diarization skipped. Transcription will continue without speaker labels.\n"
                    f"Reason: {e}\n"
                )
                diarization_turns = []

        self.root.after(
            0,
            self.status.set,
            f"Transcribing with {actual_device.upper()} ({actual_compute_type}) - 0% | ETA --:--"
        )

        try:
            segments, info = model.transcribe(
                audio_path,
                language=None,
                vad_filter=True,
                beam_size=5
            )
        except Exception as e:
            self.root.after(
                0,
                lambda: messagebox.showerror("Transcription Error", f"Gagal transcribe audio.\n\n{e}")
            )
            self.root.after(0, self.status.set, "Transcription failed")
            self.root.after(0, lambda: self.start_btn.config(state="normal"))
            return

        duration_s = float(getattr(info, "duration", 0.0)) or 0.0
        t0 = time.time()

        utterance_rows = []
        segment_ranges = []
        last_end = 0.0

        for seg in segments:
            seg_text = (seg.text or "").strip()
            st = float(seg.start)
            en = float(seg.end)

            if en > st:
                segment_ranges.append((st, en))
                last_end = max(last_end, en)

            if seg_text:
                ts = sec_to_hms(st)
                segment_speaker = assign_speaker(st, en, diarization_turns)
                self.root.after(0, self.log_live_segment, ts, segment_speaker, seg_text)

                sents = split_sentences_universal(seg_text)
                for ust, uen, s in distribute_times(st, en, sents):
                    s_clean = (s or "").strip()
                    if s_clean:
                        utterance_speaker = assign_speaker(ust, uen, diarization_turns)
                        utterance_rows.append({
                            "speaker": utterance_speaker,
                            "start_s": ust,
                            "end_s": uen,
                            "text": s_clean
                        })

            if duration_s > 0:
                progress = min(100.0, (en / duration_s) * 100.0)
                elapsed = time.time() - t0

                if progress > 0.5:
                    total_est = elapsed * (100.0 / progress)
                    eta = max(0.0, total_est - elapsed)
                    eta_str = sec_to_mmss(eta)
                else:
                    eta_str = "--:--"

                self.root.after(0, self.progress_var.set, progress)
                self.root.after(
                    0,
                    self.status.set,
                    f"Transcribing with {actual_device.upper()} ({actual_compute_type}) - {progress:.0f}% | ETA {eta_str}"
                )

        audio_duration_s = float(getattr(info, "duration", last_end))
        elapsed_min = round((time.time() - start_clock) / 60, 2)

        meta_from_gui = {
            "project": inputs["project"],
            "researcher": inputs["researcher"],
            "label": inputs["label"],
            "period": period,
        }

        try:
            save_xlsx_with_early_report(
                xlsx_path=xlsx_path,
                utterance_rows=utterance_rows,
                segment_ranges=segment_ranges,
                audio_path=audio_path,
                audio_duration_s=audio_duration_s,
                processing_minutes=elapsed_min,
                detected_language=getattr(info, "language", None),
                meta_from_gui=meta_from_gui,
                actual_device=actual_device,
                actual_compute_type=actual_compute_type,
                diarization_enabled=diarization_enabled,
                diarization_status=diarization_status,
                diarization_model=DIARIZATION_MODEL
            )
        except Exception as e:
            self.root.after(
                0,
                lambda: messagebox.showerror("Excel Error", f"Gagal simpan fail Excel.\n\n{e}")
            )
            self.root.after(0, self.status.set, "Excel save failed")
            self.root.after(0, lambda: self.start_btn.config(state="normal"))
            return

        def done_popup():
            self.progress_var.set(100.0)
            self.status.set(f"Completed - 100% | {actual_device.upper()} ({actual_compute_type})")
            self.start_btn.config(state="normal")

            messagebox.showinfo(
                "Transcription Completed Successfully",
                f"This transcription was completed by RedScribe – Research Speech Transcription System in {elapsed_min} minutes.\n\n"
                f"Whisper device used: {actual_device.upper()} ({actual_compute_type})\n"
                f"Diarization: {diarization_status}\n\n"
                "The generated output is ready for researcher-led review and verification.\n\n"
                "Feedback and suggestions for improvement are welcome at\n"
                "saprimad@moh.gov.my\n\n"
                f"Saved Excel:\n{xlsx_path}"
            )
            open_folder(out_folder)

        self.root.after(0, done_popup)


if __name__ == "__main__":
    root = tk.Tk()
    app = RedScribeApp(root)
    root.mainloop()
